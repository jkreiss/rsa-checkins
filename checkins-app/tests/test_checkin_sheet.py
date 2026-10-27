from pathlib import Path

import openpyxl
import pytest

from checkin_fixture import make_workbook
from rsa_checkins.checkin_sheet import TABS, CheckinFormatError, read_workbook

REAL = Path(__file__).resolve().parents[2] / "Football checkin sheet.xlsx"


@pytest.fixture
def rows(tmp_path):
    return read_workbook(make_workbook(tmp_path / "c.xlsx")).rows


def _find(rows, last, **match):
    return next(r for r in rows if r.values["last"] == last
                and all(r.values[k] == v for k, v in match.items()))


def test_reads_entry_rows_of_all_four_tabs(rows):
    assert [(r.tab, r.values["first"] + " " + r.values["last"]) for r in rows] == [
        ("jerseys", "Keenan Allen"), ("jerseys", "Keenan Allen"), ("jerseys", "Anthony Munoz"),
        ("jerseys", "Alvin Kamara"), ("jerseys", "Ron Jaworski"),
        ("fs-helm", "Bob Griese"),
        ("minis", "Terry Bradshaw"), ("minis", "LeGarrette Blount"),
        ("fb-balls", "Devonta Smith"), ("fb-balls", "Jason Kelce"),
    ]


def test_row_numbers_match_the_sheet(rows):
    assert rows[0].row == 3  # row 2 is the Grubauer sample
    assert _find(rows, "Griese").row == 3


def test_meanings_per_tab(rows):
    griese = _find(rows, "Griese").values
    assert (griese["rep_auth"], griese["style"], griese["model"]) == ("Authentic", "Eclipse", "Speedflex")
    bradshaw = _find(rows, "Bradshaw").values
    assert (bradshaw["league"], bradshaw["style"]) == ("NFL", "Throwback Gold")
    keenan = _find(rows, "Allen", jersey_type="Authentic").values
    assert keenan["bag"] == "Green Bag"
    assert _find(rows, "Allen", jersey_type="Custom").values["tags"] == "Colts, GREEN BAG"   # "Tag" column
    assert keenan["ws"] == "200"
    assert _find(rows, "Kelce").values["ws"] == "180"          # header "Ws "
    assert _find(rows, "Blount").values["ws"] == "54"          # header "WS price"
    assert _find(rows, "Griese").values["jersey_type"] == ""   # not a jerseys column


def test_values_are_cleaned(rows):
    blount = _find(rows, "Blount").values
    assert blount["last"] == "Blount"
    assert blount["mascot"] == ""
    assert round(float(blount["price"]), 2) == 69.99
    assert _find(rows, "Jaworski").values["mascot"] == ""
    assert _find(rows, "Allen", jersey_type="Authentic").values["inv"] == "15"


def test_rows_with_only_formula_defaults_end_the_block(tmp_path):
    path = make_workbook(tmp_path / "c.xlsx",
                         extra={"jerseys": [{"sport": "Football", "product": "Jersey", "weight": 32}]})
    assert len([r for r in read_workbook(path).rows if r.tab == "jerseys"]) == 5


def test_unknown_tabs_are_listed(tmp_path):
    path = make_workbook(tmp_path / "c.xlsx")
    wb = openpyxl.load_workbook(path)
    wb.create_sheet("notes")
    wb.save(path)
    assert read_workbook(path).ignored_tabs == ["notes"]


def test_not_a_checkin_sheet(tmp_path):
    path = tmp_path / "other.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.save(path)
    with pytest.raises(CheckinFormatError) as e:
        read_workbook(path)
    assert all(t in str(e.value) for t in TABS)


def test_missing_required_header(tmp_path):
    path = make_workbook(tmp_path / "c.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["jerseys"]
    for cell in ws[1]:
        if cell.value == "cert":
            cell.value = "certificate"
    wb.save(path)
    with pytest.raises(CheckinFormatError) as e:
        read_workbook(path)
    assert "jerseys" in str(e.value) and "cert" in str(e.value)


def test_real_workbook_smoke():
    if not REAL.exists():
        pytest.skip("real check-in workbook not present")
    read = read_workbook(REAL)
    assert {r.tab for r in read.rows} == set(TABS)
    keenan = next(r for r in read.rows if r.values["last"] == "Allen" and r.values["jersey_type"] == "Authentic")
    assert keenan.values["model"] == "Mitchell & Ness"
    assert all(r.values["ws"].lower() != "x" and r.values["date"].lower() != "sample" for r in read.rows)
